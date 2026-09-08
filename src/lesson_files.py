"""Local lesson uploads using the same parsers as document ingestion; no OCR or LLM calls."""

import io
import re
import zipfile
from pathlib import Path

import fitz
from docx import Document

MAX_LESSON_FILE_BYTES = 10 * 1024 * 1024
NO_TEXT_ERROR = "未检测到可提取文本，请粘贴课程内容或上传可选择文字的文件"


def parse_lesson_file(upload):
    """Validate and extract entirely before writing a persistent upload."""
    filename = upload.filename
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".docx"}:
        raise ValueError("Only PDF and DOCX files are supported.")
    data = upload.stream.read(MAX_LESSON_FILE_BYTES + 1)
    if len(data) > MAX_LESSON_FILE_BYTES:
        raise ValueError("Lesson files must be 10 MB or smaller.")
    if not data:
        raise ValueError("The uploaded file is empty. Choose a file containing lesson text.")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("The file is not a valid PDF. Changing its extension does not convert it.")
    if extension == ".docx" and not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("The file is not a valid DOCX. Changing its extension does not convert it.")
    try:
        if extension == ".pdf":
            with fitz.open(stream=data, filetype="pdf") as doc:
                if not doc.is_pdf or doc.is_repaired:
                    raise ValueError("The PDF is damaged. Export it again and retry.")
                if doc.needs_pass:
                    raise ValueError("Password-protected PDFs are not supported. Upload an unlocked copy.")
                text = "\n".join(page.get_text() for page in doc).strip()
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                # DOCX is a ZIP package. Bound expanded content before python-docx reads it.
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("The DOCX expands beyond the 50 MB processing limit.")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("The file is not a valid DOCX document.")
                if archive.testzip() is not None:
                    raise ValueError("The DOCX is damaged. Export it again and retry.")
            doc = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in doc.paragraphs]
            parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            text = "\n".join(parts).strip()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The file is damaged or could not be parsed. Export a new PDF or DOCX and retry.") from exc
    if not text:
        raise ValueError(NO_TEXT_ERROR)
    return data, extension, text


def lesson_file_path(directory, stored_filename):
    """Accept only generated basenames and reject symlinks escaping private storage."""
    if not stored_filename or not re.fullmatch(r"[a-f0-9]{32}\.(pdf|docx)", stored_filename):
        raise ValueError("Invalid stored filename")
    root = Path(directory).resolve()
    path = (root / stored_filename).resolve()
    if path.parent != root:
        raise ValueError("Invalid file location")
    return path
