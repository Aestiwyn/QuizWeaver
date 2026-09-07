"""Lesson recording, persistence, document validation and class isolation."""

import io
import re
import zipfile
from datetime import date
from unittest.mock import patch

import fitz
import pytest
from docx import Document

from src.database import LessonLog, get_session


def pdf_bytes(text="Photosynthesis from the PDF"):
    with fitz.open() as doc:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
        return doc.tobytes()


def docx_bytes():
    doc = Document()
    doc.add_paragraph("Cell division from the DOCX")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table material"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture(autouse=True)
def upload_folder(flask_app, tmp_path):
    folder = tmp_path / "lesson_uploads"
    flask_app.config["LESSON_UPLOAD_DIR"] = str(folder)
    return folder


def saved_lessons(app):
    with get_session(app.config["DB_ENGINE"]) as session:
        return session.query(LessonLog).order_by(LessonLog.id).all()


def test_text_date_detail_and_refresh(flask_client, flask_app):
    response = flask_client.post("/classes/1/lessons/new", data={
        "content": "Manual lesson", "lesson_date": "2024-02-29", "topics": "Cells", "notes": "Review later",
    })
    assert response.status_code == 303
    lesson = saved_lessons(flask_app)[-1]
    assert lesson.date == date(2024, 2, 29)
    assert response.location.endswith(f"/classes/1/lessons/{lesson.id}")
    for _ in range(2):
        html = flask_client.get(response.location).get_data(as_text=True)
        for value in ["2024-02-29", "Manual lesson", "Cells", "Review later", "Test Class"]:
            assert value in html
        assert f"/classes/1/generate?lesson_id={lesson.id}" in html
        assert "Generate Quiz from This Lesson" in html
    assert b"2024-02-29" in flask_client.get("/classes/1/lessons").data


@pytest.mark.parametrize("invalid", ["2024-02-30", "not-a-date", "20240906", "2024-1-2"])
def test_invalid_date_preserves_form(flask_client, flask_app, upload_folder, invalid):
    response = flask_client.post("/classes/1/lessons/new", data={
        "content": "Keep my content", "lesson_date": invalid, "topics": "Keep topics", "notes": "Keep notes",
        "lesson_file": (io.BytesIO(pdf_bytes()), "lesson.pdf"),
    })
    assert response.status_code == 400
    html = response.get_data(as_text=True)
    for value in [invalid, "Keep my content", "Keep topics", "Keep notes", "YYYY-MM-DD", "select the file again"]:
        assert value in html
    assert not saved_lessons(flask_app)
    assert not list(upload_folder.glob("*"))


def test_missing_date_defaults_today(flask_client, flask_app):
    assert flask_client.post("/classes/1/lessons/new", data={"content": "A lesson"}).status_code == 303
    assert saved_lessons(flask_app)[0].date == date.today()


@pytest.mark.parametrize("kind", ["pdf", "docx"])
@pytest.mark.parametrize("manual", ["", "Manual source"])
def test_files_and_combined_content(flask_client, flask_app, upload_folder, kind, manual):
    raw = pdf_bytes() if kind == "pdf" else docx_bytes()
    original = f"课程.{kind}"
    response = flask_client.post("/classes/1/lessons/new", data={
        "content": manual, "lesson_file": (io.BytesIO(raw), original),
    })
    assert response.status_code == 303
    lesson = saved_lessons(flask_app)[0]
    assert lesson.content == manual
    assert f"from the {kind.upper()}" in lesson.extracted_text
    if kind == "docx":
        assert "Table material" in lesson.extracted_text
    assert lesson.original_filename == original
    assert re.fullmatch(r"[a-f0-9]{32}\.(pdf|docx)", lesson.stored_filename)
    assert (upload_folder / lesson.stored_filename).read_bytes() == raw
    assert manual in lesson.generation_content
    assert lesson.extracted_text in lesson.generation_content
    html = flask_client.get(response.location).get_data(as_text=True)
    assert original in html and "Manual Lesson Content" in html and "Extracted File Content" in html
    download = flask_client.get(response.location + "/download")
    assert download.status_code == 200 and download.data == raw
    assert "attachment;" in download.headers["Content-Disposition"]
    download.close()


@pytest.mark.parametrize("manual", ["", "Keep manual text"])
@pytest.mark.parametrize("filename,raw,error", [
    ("bad.txt", b"text", "Only PDF and DOCX"),
    ("empty.pdf", b"", "empty"),
    ("fake.pdf", b"not PDF", "valid PDF"),
    ("broken.pdf", b"%PDF-1.7\ncorrupt", "damaged"),
    ("fake.docx", b"not DOCX", "valid DOCX"),
    ("blank.pdf", pdf_bytes(""), "未检测到可提取文本，请粘贴课程内容或上传可选择文字的文件"),
    ("big.pdf", b"x" * (10 * 1024 * 1024 + 1), "10 MB"),
], ids=["extension", "empty", "fake-pdf", "damaged-pdf", "fake-docx", "no-text", "oversize"])
def test_invalid_upload_never_saves(flask_client, flask_app, upload_folder, filename, raw, error, manual):
    response = flask_client.post("/classes/1/lessons/new", data={
        "content": manual, "topics": "Keep topics", "notes": "Keep notes", "lesson_date": "2024-03-05",
        "lesson_file": (io.BytesIO(raw), filename),
    })
    assert response.status_code == 400
    html = response.get_data(as_text=True)
    for value in [manual, "Keep topics", "Keep notes", "2024-03-05", error, "select the file again"]:
        assert value in html
    assert "Lesson logged successfully" not in html
    assert not saved_lessons(flask_app)
    assert not list(upload_folder.glob("*"))


@pytest.mark.parametrize("content", ["", " \n\t "])
def test_empty_submission(flask_client, flask_app, content):
    response = flask_client.post("/classes/1/lessons/new", data={"content": content})
    assert response.status_code == 400
    assert b"Enter lesson content or upload" in response.data
    assert not saved_lessons(flask_app)


def test_download_authorization_and_legacy(flask_client, flask_app, upload_folder):
    response = flask_client.post("/classes/1/lessons/new", data={"lesson_file": (io.BytesIO(pdf_bytes()), "lesson.pdf")})
    url = response.location
    assert flask_client.get(url.replace("/classes/1/", "/classes/2/")).status_code == 404
    assert flask_client.get(url.replace("/classes/1/", "/classes/2/") + "/download").status_code == 404
    assert flask_client.get("/classes/999/lessons/1/download").status_code == 404
    assert flask_client.get("/classes/1/lessons/999/download").status_code == 404
    with flask_app.test_client() as anonymous:
        assert anonymous.get(url + "/download").status_code == 303
        assert anonymous.get(url).status_code == 303
    legacy = flask_client.post("/classes/1/lessons/new", data={"content": "Old text-only lesson"}).location
    assert flask_client.get(legacy).status_code == 200
    assert flask_client.get(legacy + "/download").status_code == 404


def test_storage_names_and_path_tampering(flask_client, flask_app, upload_folder):
    for _ in range(2):
        assert flask_client.post("/classes/1/lessons/new", data={
            "lesson_file": (io.BytesIO(pdf_bytes()), "../../outside.pdf"),
        }).status_code == 303
    lessons = saved_lessons(flask_app)
    assert lessons[0].stored_filename != lessons[1].stored_filename
    assert len(list(upload_folder.glob("*"))) == 2
    with get_session(flask_app.config["DB_ENGINE"]) as session:
        lesson = session.get(LessonLog, lessons[0].id)
        lesson.stored_filename = "../outside.pdf"
        session.commit()
    assert flask_client.get(f"/classes/1/lessons/{lessons[0].id}/download").status_code == 404


def test_database_failure_removes_upload(flask_client, flask_app, upload_folder):
    with patch("sqlalchemy.orm.Session.commit", side_effect=RuntimeError("Database unavailable")):
        response = flask_client.post("/classes/1/lessons/new", data={
            "content": "Preserve me", "lesson_file": (io.BytesIO(pdf_bytes()), "lesson.pdf"),
        })
    assert response.status_code == 500
    assert b"Preserve me" in response.data
    assert not saved_lessons(flask_app)
    assert not list(upload_folder.glob("*"))


def test_generation_uses_both_sources_and_locks_class(flask_client, flask_app):
    response = flask_client.post("/classes/1/lessons/new", data={
        "content": "Manual instructions", "lesson_date": "2020-01-01",
        "lesson_file": (io.BytesIO(pdf_bytes()), "source.pdf"),
    })
    lesson = saved_lessons(flask_app)[0]
    url = f"/classes/1/generate?lesson_id={lesson.id}"
    page = flask_client.get(url)
    assert page.status_code == 200
    assert f'<option value="{lesson.id}" selected>'.encode() in page.data
    assert b'<select name="class_id"' not in page.data
    with patch("src.web.blueprints.quizzes.generate_quiz", return_value=None) as generate:
        flask_client.post(
            "/classes/1/generate",
            data={"source_mode": "recorded_lesson", "lesson_id": str(lesson.id), "content_text": "Ignored extra"},
        )
        context = generate.call_args.kwargs["content_text"]
        assert "Manual instructions" in context and "Photosynthesis from the PDF" in context
        assert "Ignored extra" not in context
        assert generate.call_args.kwargs["class_id"] == 1
    assert flask_client.get(f"/classes/2/generate?lesson_id={lesson.id}").status_code == 404
    assert flask_client.post(
        "/classes/2/generate", data={"source_mode": "recorded_lesson", "lesson_id": lesson.id}
    ).status_code == 404


def test_unexpected_parser_error_preserves_text(flask_client, flask_app, upload_folder):
    with patch("src.web.blueprints.classes.parse_lesson_file", side_effect=RuntimeError("Parser crashed")):
        response = flask_client.post("/classes/1/lessons/new", data={
            "content": "Keep this", "lesson_file": (io.BytesIO(pdf_bytes()), "source.pdf"),
        })
    assert response.status_code == 400
    assert b"Keep this" in response.data and b"could not be parsed" in response.data
    assert not saved_lessons(flask_app) and not list(upload_folder.glob("*"))


def test_ten_mb_file_and_real_csrf(flask_client, flask_app):
    flask_app.config["WTF_CSRF_ENABLED"] = True
    form = flask_client.get("/classes/1/lessons/new").get_data(as_text=True)
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', form).group(1)
    stream = io.BytesIO(docx_bytes())
    padding_name = "padding.bin"
    padding_size = 10 * 1024 * 1024 - len(stream.getvalue()) - 76 - 2 * len(padding_name)
    with zipfile.ZipFile(stream, "a") as archive:
        archive.writestr(padding_name, b"x" * padding_size)
    assert len(stream.getvalue()) == 10 * 1024 * 1024
    stream.seek(0)
    response = flask_client.post("/classes/1/lessons/new", data={
        "csrf_token": csrf, "lesson_file": (stream, "large.docx"),
    })
    assert response.status_code == 303
    assert "Cell division" in saved_lessons(flask_app)[0].extracted_text


def test_entire_request_limit(flask_client, flask_app, upload_folder):
    response = flask_client.post("/classes/1/lessons/new", data={
        "lesson_file": (io.BytesIO(b"x" * (12 * 1024 * 1024)), "huge.pdf"),
    })
    assert response.status_code == 413
    assert b"10 MB" in response.data and b"Back button" in response.data
    assert not saved_lessons(flask_app) and not list(upload_folder.glob("*"))


def test_delete_is_class_scoped_and_removes_original(flask_client, flask_app, upload_folder):
    response = flask_client.post("/classes/1/lessons/new", data={"lesson_file": (io.BytesIO(pdf_bytes()), "source.pdf")})
    lesson = saved_lessons(flask_app)[0]
    assert flask_client.post(f"/classes/2/lessons/{lesson.id}/delete").status_code == 404
    assert (upload_folder / lesson.stored_filename).exists()
    assert flask_client.post(response.location + "/delete").status_code == 303
    assert not saved_lessons(flask_app) and not list(upload_folder.glob("*"))


@pytest.mark.parametrize("kind", ["renamed-pdf", "renamed-docx", "zip-only", "broken-docx", "encrypted-pdf", "image-pdf", "empty-docx"])
def test_real_format_validation(flask_client, flask_app, upload_folder, kind):
    filename = "source.docx"
    if kind == "renamed-pdf":
        data = pdf_bytes()
    elif kind == "renamed-docx":
        filename, data = "source.pdf", docx_bytes()
    elif kind in {"zip-only", "broken-docx"}:
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("word/document.xml" if kind == "broken-docx" else "readme.txt", "not an Office document")
        data = stream.getvalue()
    elif kind == "empty-docx":
        stream = io.BytesIO()
        Document().save(stream)
        data = stream.getvalue()
    else:
        filename = "source.pdf"
        with fitz.open() as doc:
            page = doc.new_page()
            if kind == "encrypted-pdf":
                page.insert_text((72, 72), "Locked text")
                data = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="reader")
            else:
                # Actual image content, not an empty PDF page.
                from PIL import Image

                stream = io.BytesIO()
                Image.new("RGB", (20, 20), "blue").save(stream, format="PNG")
                page.insert_image(fitz.Rect(72, 72, 172, 172), stream=stream.getvalue())
                data = doc.tobytes()
    response = flask_client.post("/classes/1/lessons/new", data={"lesson_file": (io.BytesIO(data), filename)})
    assert response.status_code == 400
    html = response.get_data(as_text=True)
    assert "select the file again" in html
    if kind in {"image-pdf", "empty-docx"}:
        assert "未检测到可提取文本，请粘贴课程内容或上传可选择文字的文件" in html
    assert not saved_lessons(flask_app) and not list(upload_folder.glob("*"))
